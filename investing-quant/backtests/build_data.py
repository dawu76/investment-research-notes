# Run from this directory:  uv run --with xlrd --with openpyxl build_data.py
"""Build data/annual_returns.csv and data/monthly_returns.csv from data/raw/ (no network access)."""

import csv
import html
import json
import io
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import xlrd

HERE = Path(__file__).resolve().parent
RAW = HERE / "data" / "raw"
OUT = HERE / "data"

FIRST_YEAR = 1971
LAST_YEAR = 2025
FIRST_MONTH = "1971-01"

COLUMNS = ["us_stocks", "intl_stocks", "us_small_value", "tsy_10y", "tbill", "gold", "cpi"]
# Layer-2 trend sources. Blank before each series starts.
TREND_COLUMNS = ["cta_index", "aqr_tsmom"]
# Layer-3 TIPS series: modeled from market real yields, synthetic from estimated real yields, and an actual fund.
TIPS_COLUMNS = ["tips_10y", "tips_synth", "tips_fund"]


def damodaran_annual():
    book = xlrd.open_workbook(str(RAW / "histretSP.xls"))
    sheet = book.sheet_by_name("Returns by year")
    header_row = next(r for r in range(sheet.nrows) if sheet.cell_value(r, 0) == "Year")
    headers = [str(sheet.cell_value(header_row, c)) for c in range(sheet.ncols)]

    def col(prefix):
        return next(i for i, h in enumerate(headers) if h.startswith(prefix))

    idx = {
        "us_stocks": col("S&P 500"),
        "tbill": col("3-month T.Bill"),
        "tsy_10y": col("US T. Bond"),
        "gold": col("Gold"),
    }
    out = {}
    for r in range(header_row + 1, sheet.nrows):
        year = sheet.cell_value(r, 0)
        if not isinstance(year, float):
            break
        out[int(year)] = {k: float(sheet.cell_value(r, c)) for k, c in idx.items()}
    return out


def novelinvestor_eafe():
    text = (RAW / "novelinvestor_historical_returns.html").read_text(encoding="utf-8", errors="replace")
    table = re.findall(r"<table.*?</table>", text, re.S)[0]
    rows = re.findall(r"<tr.*?</tr>", table, re.S)

    def cells(row):
        return [html.unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", row, re.S)]

    header = cells(rows[0])
    col = header.index("International Stocks")
    out = {}
    for row in rows[1:]:
        c = cells(row)
        if c[0].isdigit() and c[col]:
            out[int(c[0])] = float(c[col].rstrip("%")) / 100
    return out


def french_sections(zip_name):
    with zipfile.ZipFile(RAW / zip_name) as z:
        text = z.read(z.namelist()[0]).decode("latin-1")
    sections = {}
    current = None
    for line in text.splitlines():
        stripped = line.strip()
        first = stripped.split(",")[0].strip()
        if first.isdigit():
            if current is not None:
                sections[current].append([p.strip() for p in stripped.split(",")])
        elif stripped and not stripped.startswith(","):
            current = stripped
            sections[current] = []
        elif stripped.startswith(","):
            if current is None:
                current = "header"
                sections[current] = []
            sections[current + "::cols"] = [p.strip() for p in stripped.split(",")]
    return sections


def french_small_value():
    sec = french_sections("6_Portfolios_2x3_CSV.zip")
    cols = sec["Average Value Weighted Returns -- Monthly::cols"]
    i = cols.index("SMALL HiBM")
    monthly = {f"{r[0][:4]}-{r[0][4:]}": float(r[i]) / 100 for r in sec["Average Value Weighted Returns -- Monthly"]}
    annual = {int(r[0]): float(r[i]) / 100 for r in sec["Average Value Weighted Returns -- Annual"]}
    return annual, monthly


def french_market_monthly():
    sec = french_sections("F-F_Research_Data_Factors_CSV.zip")
    key = next(k for k in sec if not k.endswith("::cols") and sec[k] and len(sec[k][0][0]) == 6)
    cols = sec[key + "::cols"]
    mkt, rf = cols.index("Mkt-RF"), cols.index("RF")
    return {f"{r[0][:4]}-{r[0][4:]}": (float(r[mkt]) + float(r[rf])) / 100 for r in sec[key]}


def fred(series):
    with open(RAW / f"fred_{series}.csv") as f:
        return {row["observation_date"][:7]: float(row[series]) for row in csv.DictReader(f) if row[series] not in ("", ".")}


def fred_daily(series):
    with open(RAW / f"fred_{series}.csv") as f:
        return {row["observation_date"]: float(row[series]) for row in csv.DictReader(f) if row[series] not in ("", ".")}


def month_end(daily):
    """Last observation in each month. Keys go from YYYY-MM-DD to YYYY-MM."""
    return {day[:7]: daily[day] for day in sorted(daily)}


def lbma_gold_daily():
    rows = json.loads((RAW / "lbma_gold_pm.json").read_text())
    return {r["d"]: r["v"][0] for r in rows if r["v"][0]}


def barclay_cta_annual():
    """Barclay CTA Index annual returns (net of fees). Negative years are wrapped in a red <font> tag."""
    text = (RAW / "barclay_cta_index.html").read_text(encoding="utf-8", errors="replace")
    rows = re.findall(r"<tr><td>(\d{4})</td><td>(?:<font[^>]*>)?(-?[\d.]+)(?:</font>)?</td></tr>", text)
    return {int(y): float(v) / 100 for y, v in rows if int(y) <= LAST_YEAR}


def aqr_tsmom_monthly_excess():
    """AQR TSMOM factor (all assets), monthly excess return over T-bills, gross of fees."""
    ws = openpyxl.load_workbook(RAW / "aqr_tsmom.xlsx", read_only=True, data_only=True)["TSMOM Factors"]
    rows = list(ws.iter_rows(values_only=True))
    col = next(r for r in rows if r[1] == "TSMOM").index("TSMOM")
    return {r[0].strftime("%Y-%m"): float(r[col]) for r in rows if hasattr(r[0], "strftime") and r[col] is not None}


def gold_monthly_avg():
    with open(RAW / "gold_monthly.csv") as f:
        return {row["Date"][:7]: float(row["Price"]) for row in csv.DictReader(f)}


def bond_price(coupon_pct, yield_pct, years):
    """Price per 100 face of a semiannual-coupon bond."""
    y = yield_pct / 200
    n = years * 2
    c = coupon_pct / 2
    if y == 0:
        return c * n + 100
    return c * (1 - (1 + y) ** -n) / y + 100 * (1 + y) ** -n


def tsy10_monthly_tr(gs10):
    """Constant-maturity 10y total return: buy a par bond at last month's yield, reprice a month later at this month's yield."""
    months = sorted(gs10)
    out = {}
    for prev, cur in zip(months, months[1:]):
        y0, y1 = gs10[prev], gs10[cur]
        price = bond_price(y0, y1, 10 - 1 / 12)
        out[cur] = (price + y0 / 12) / 100 - 1
    return out


def prev_month(m):
    year, month = int(m[:4]), int(m[5:])
    return f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"


def next_month(m):
    year, month = int(m[:4]), int(m[5:])
    return f"{year + 1}-01" if month == 12 else f"{year}-{month + 1:02d}"


def fill_single_gaps(levels):
    """Geometric midpoint for one missing month between two present months (BLS never published Oct 2025 CPI)."""
    out = dict(levels)
    for m in sorted(levels):
        gap, after = next_month(m), next_month(next_month(m))
        if gap not in levels and after in levels:
            out[gap] = (levels[m] * levels[after]) ** 0.5
    return out


def tips_monthly_tr(real_yields, cpi_m):
    """Constant-maturity 10y TIPS: par real bond at last month's real yield, repriced at this month's, principal times CPI."""
    out = {}
    for cur, y1 in real_yields.items():
        prev = prev_month(cur)
        if prev in real_yields and cur in cpi_m:
            y0 = real_yields[prev]
            out[cur] = (bond_price(y0, y1, 10 - 1 / 12) + y0 / 12) / 100 * (1 + cpi_m[cur]) - 1
    return out


def gsw_tips_par10_daily():
    with open(RAW / "feds200805_tipspy10.csv") as f:
        return {r["Date"]: float(r["TIPSPY10"]) for r in csv.DictReader(f)}


def livingston_expected_cpi():
    """Median CPI inflation forecast (% a year, base period to 12 months ahead), keyed by June/December survey month."""
    ws = openpyxl.load_workbook(RAW / "livingston_median_growth.xlsx", read_only=True, data_only=True)["CPI"]
    rows = list(ws.iter_rows(values_only=True))
    col = rows[0].index("G_BP_To_12M")
    return {r[0].strftime("%Y-%m"): float(r[col]) for r in rows[1:]
            if hasattr(r[0], "strftime") and isinstance(r[col], (int, float))}


def synthetic_real_yields(nominal):
    """Month-end nominal 10y yield minus expected inflation: Livingston (carried forward) before 1982, Cleveland Fed after."""
    liv, cle = livingston_expected_cpi(), fred("EXPINF10YR")
    out, expected = {}, None
    for m in sorted(nominal):
        expected = cle.get(m) if m >= "1982-01" else liv.get(m, expected)
        if expected is not None:
            out[m] = nominal[m] - expected
    return out


def yahoo_monthly_tr(ticker):
    """Month-over-month change in Yahoo's dividend-adjusted close."""
    r = json.loads((RAW / f"yahoo_{ticker}.json").read_text())["chart"]["result"][0]
    months = [datetime.fromtimestamp(t, timezone.utc).strftime("%Y-%m") for t in r["timestamp"]]
    return pct_change({m: v for m, v in zip(months, r["indicators"]["adjclose"][0]["adjclose"]) if v is not None})


def compound_year(monthly, year):
    months = [f"{year}-{m:02d}" for m in range(1, 13)]
    if not all(m in monthly for m in months):
        return None
    return math.prod(1 + monthly[m] for m in months) - 1


def fmt(v):
    return f"{v:.6f}" if isinstance(v, float) else ("" if v is None else v)


def pct_change(levels):
    """Month-over-month change. Skipped when the prior calendar month is missing (e.g. no Oct 2025 CPI)."""
    return {m: levels[m] / levels[prev_month(m)] - 1 for m in sorted(levels) if prev_month(m) in levels}


def main():
    dam = damodaran_annual()
    eafe = novelinvestor_eafe()
    sv_annual, sv_monthly = french_small_value()
    cpi_levels = fred("CPIAUCNS")
    gold_levels = gold_monthly_avg()
    tb3 = fred("TB3MS")
    cta = barclay_cta_annual()
    aqr = {m: x + tb3[m] / 1200 for m, x in aqr_tsmom_monthly_excess().items()}
    dgs10 = month_end(fred_daily("DGS10"))
    cpi_filled = pct_change(fill_single_gaps(cpi_levels))
    tips = {
        "tips_10y": tips_monthly_tr(month_end(gsw_tips_par10_daily()), cpi_filled),
        "tips_synth": tips_monthly_tr(synthetic_real_yields(dgs10), cpi_filled),
        "tips_fund": yahoo_monthly_tr("VIPSX"),
    }

    annual_rows = []
    for year in range(FIRST_YEAR, LAST_YEAR + 1):
        cpi = cpi_levels[f"{year}-12"] / cpi_levels[f"{year - 1}-12"] - 1
        gold_avg = sum(gold_levels[f"{year}-{m:02d}"] for m in range(1, 13)) / 12
        gold_avg_prev = sum(gold_levels[f"{year - 1}-{m:02d}"] for m in range(1, 13)) / 12
        annual_rows.append({
            "year": year,
            "us_stocks": dam[year]["us_stocks"],
            "intl_stocks": eafe[year],
            "us_small_value": sv_annual[year],
            "tsy_10y": dam[year]["tsy_10y"],
            "tbill": dam[year]["tbill"],
            "gold": dam[year]["gold"],
            "cpi": cpi,
            "gold_annual_avg": gold_avg / gold_avg_prev - 1,
            "cta_index": cta.get(year),
            "aqr_tsmom": compound_year(aqr, year),
            **{col: compound_year(series, year) for col, series in tips.items()},
        })

    with open(OUT / "annual_returns.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["year"] + COLUMNS + ["gold_annual_avg"] + TREND_COLUMNS + TIPS_COLUMNS)
        w.writeheader()
        for row in annual_rows:
            w.writerow({k: fmt(v) for k, v in row.items()})

    mkt = french_market_monthly()
    tsy = tsy10_monthly_tr(dgs10)
    cpi_m = pct_change(cpi_levels)
    gold_m = pct_change(month_end(lbma_gold_daily()))
    months = sorted(m for m in mkt if m >= FIRST_MONTH and all(m in s for s in (tsy, tb3, gold_m, sv_monthly)))

    with open(OUT / "monthly_returns.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["month"] + COLUMNS + ["aqr_tsmom"] + TIPS_COLUMNS)
        for m in months:
            w.writerow([m, f"{mkt[m]:.6f}", "", f"{sv_monthly[m]:.6f}", f"{tsy[m]:.6f}",
                        f"{tb3[m] / 1200:.6f}", f"{gold_m[m]:.6f}", fmt(cpi_m.get(m)), fmt(aqr.get(m))]
                       + [fmt(tips[col].get(m)) for col in TIPS_COLUMNS])

    print(f"annual: {len(annual_rows)} rows {FIRST_YEAR}-{LAST_YEAR}; monthly: {len(months)} rows {months[0]}..{months[-1]}")


if __name__ == "__main__":
    main()
